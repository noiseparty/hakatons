"""Banku bankomātu saraksts (Finance Latvia / bankas, 22.09.2026, hakatonam) → CSV kartei.

Avots: Kritisko_ATM saraksts_22.09.2026_hakatonam.xlsx (15 MB, repozitorijā NAV — licence nav norādīta, skat.
atseviski_dati/README.md). Lapa "Visi ATM": ATM ID, ADRESE, LONG, LAT (kolonnas samainītas: LONG ir platums ~56,9),
pieejamība (IEROB / NEIEROB = 24/7), veids (CRM = iemaksas un izmaksas, COM = tikai izmaksas), kritiskuma pazīme (KRIT),
MFI kods (900 SEB, 814 Luminor, 500 Citadele, 717 Swedbank). Lapa "Tikai kritiskie_ATM" — tās pašas KRIT rindas.

Izvade (src/karte/dati/):
  bankomati_bankas.csv  viens bankomāts rindā (845): id, adrese, lat, lon, pieejamiba, veids, kritiskais, banka
  bankomati_vietas.csv  viena vieta rindā (vairāki bankomāti vienā adresē → viens punkts kartē) — to ielādē ielade.py

  uv run --no-project --python 3.12 --with openpyxl src/karte/db/bankomati.py "C:/…/Kritisko_ATM saraksts_22.09.2026_hakatonam.xlsx"
  python3 src/karte/db/bankomati.py --tikai-vietas      # tikai no bankomati_bankas.csv (bez xlsx; VPS)
"""

import argparse
import csv
import pathlib
from collections import OrderedDict

DATI = pathlib.Path(__file__).resolve().parents[1] / "dati"
BANKAS = {"900": "SEB", "814": "Luminor", "500": "Citadele", "717": "Swedbank"}
LAUKI = ["id", "adrese", "lat", "lon", "pieejamiba", "veids", "kritiskais", "banka"]


def no_xlsx(cels):
    import openpyxl
    wb = openpyxl.load_workbook(cels, read_only=True, data_only=True)
    kritiskie = set()
    for r in wb["Tikai kritiskie_ATM"].iter_rows(values_only=True):  # galvene var nebūt 1. rindā — pēc satura
        if r and r[0] and len(r) > 6 and str(r[6] or "").strip().upper() == "KRIT":
            kritiskie.add(str(r[0]).strip())
    rindas = []
    for r in wb["Visi ATM"].iter_rows(max_col=8, values_only=True):
        if not r or not r[0] or str(r[0]).strip().upper().replace("_", " ") == "ATM ID":
            continue
        atm, adrese, long_, lat_, pieejamiba, veids, krit, mfi = r
        lat, lon = float(str(long_).replace(",", ".")), float(str(lat_).replace(",", "."))  # kolonnas samainītas
        if 20 < lat < 29 and 55 < lon < 59:  # ja kādā rindā tomēr pareizā secībā
            lat, lon = lon, lat
        if not (55 < lat < 59 and 20 < lon < 29):
            continue
        atm = str(atm).strip()
        rindas.append({"id": atm, "adrese": str(adrese or "").strip(), "lat": f"{lat:.6f}", "lon": f"{lon:.6f}",
                       "pieejamiba": str(pieejamiba or "").strip().upper(), "veids": str(veids or "").strip().upper(),
                       "kritiskais": "1" if str(krit or "").strip().upper() == "KRIT" or atm in kritiskie else "0",
                       "banka": BANKAS.get(str(mfi).strip(), str(mfi or "").strip())})
    rindas.sort(key=lambda x: x["id"])
    with open(DATI / "bankomati_bankas.csv", "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=LAUKI, lineterminator="\n")
        w.writeheader()
        w.writerows(rindas)
    print(f"bankomati_bankas.csv: {len(rindas)} bankomāti, kritiskie {sum(r['kritiskais'] == '1' for r in rindas)}")


def vietas():
    """Bankomāti tajā pašā punktā (≤ ~1 m) → viena vieta: skaits, bankas, vai ir iemaksas, 24/7, kritiskais."""
    grupas = OrderedDict()
    with open(DATI / "bankomati_bankas.csv", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            grupas.setdefault((round(float(r["lat"]), 5), round(float(r["lon"]), 5)), []).append(r)
    rez = []
    for (lat, lon), g in grupas.items():
        bankas = list(OrderedDict.fromkeys(r["banka"] for r in g))
        krit = any(r["kritiskais"] == "1" for r in g)
        rez.append({
            "id": min(r["id"] for r in g),
            "nosaukums": ("Kritiskais bankomāts" if krit else "Bankomāts") + f" · {', '.join(bankas)}" + (f" ({len(g)})" if len(g) > 1 else ""),
            "adrese": g[0]["adrese"], "lat": lat, "lon": lon, "bankas": ", ".join(bankas), "skaits": len(g),
            "iemaksas": "1" if any(r["veids"] == "CRM" for r in g) else "0",
            "pieejamiba_24h": "1" if any(r["pieejamiba"] == "NEIEROB" for r in g) else "0",
            "kritiskais": "1" if krit else "0", "atm_id": " ".join(sorted(r["id"] for r in g)),
        })
    with open(DATI / "bankomati_vietas.csv", "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rez[0]), lineterminator="\n")
        w.writeheader()
        w.writerows(rez)
    print(f"bankomati_vietas.csv: {len(rez)} vietas, kritiskās {sum(r['kritiskais'] == '1' for r in rez)}")


if __name__ == "__main__":
    a = argparse.ArgumentParser()
    a.add_argument("xlsx", nargs="?")
    a.add_argument("--tikai-vietas", action="store_true")
    args = a.parse_args()
    if not args.tikai_vietas:
        if not args.xlsx:
            a.error("vajag xlsx ceļu (vai --tikai-vietas)")
        no_xlsx(args.xlsx)
    vietas()
